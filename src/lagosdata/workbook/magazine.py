"""Configurable magazine XLSX, retaining the reference workbook's visual language."""
from collections import Counter
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from .style import *
from ..tree import TREE

COLS=['ID','Business Name','Zone','Catchment','Street / Cluster','Category Group','Subcategory','Source Category','Legacy 18-Category','Full Address','Phone','Website','Google Maps','Google Rating','Google Reviews','Has Phone','Latitude','Longitude','Data Source','Date Added','Last Checked','Verification','Prospect Priority','Suggested Package','Sales Status','Sales Notes']
def build_magazine(records,config,searches,output):
    records=list(records); wb=Workbook(); ws=wb.active; ws.title='MASTER DATABASE'; ws.append(COLS)
    for i,r in enumerate(records,1):
        ws.append([i,r.get('name',''),r.get('zone',''),r.get('area',''),r.get('street',''),r.get('group',''),r.get('sub',''),r.get('label',''),r.get('legacy18',''),r.get('addr',''),r.get('phone',''),r.get('website',''),r.get('maps',''),r.get('rating') or '',r.get('reviews') or '', 'Yes' if r.get('phone') else 'No',r.get('lat') or '',r.get('lng') or '',r.get('source',''),r.get('added',''),'',r.get('verification',''),r.get('priority',''),r.get('package',''),r.get('sales_status','Not Contacted'),r.get('notes','')])
    style_header(ws); finish(ws,[6,38,18,28,26,30,30,28,20,44,18,30,34,12,12,12,14,14,28,14,14,22,32,16,18,32],'C2')
    last=len(records)+1; dv=DataValidation(type='list',formula1='"Not Contacted,Contacted,Meeting Booked,Proposal Sent,Won,Declined,Do Not Contact"');ws.add_data_validation(dv);dv.add(f'Y2:Y{last}')
    dv2=DataValidation(type='list',formula1='"Premium,Standard,Listing"');ws.add_data_validation(dv2);dv2.add(f'X2:X{last}')
    dash=wb.create_sheet('DASHBOARD',0);dash.append(['LAGOS COMMUNITY BUSINESS DATABASE']);dash.append(['KPI','Value']);style_header(dash,2)
    for label,formula in [('Total businesses',f'=COUNTA(\'MASTER DATABASE\'!B2:B{last})'),('With phone',f'=COUNTIF(\'MASTER DATABASE\'!P2:P{last},"Yes")'),('Core catchment',f'=COUNTIF(\'MASTER DATABASE\'!C2:C{last},"Core catchment")'),('Ring areas',f'=COUNTIF(\'MASTER DATABASE\'!C2:C{last},"Ring area")')]:dash.append([label,formula])
    readme=wb.create_sheet('README');readme.append(['ABOUT THIS DATABASE']);readme.append(['Source attribution','© OpenStreetMap contributors']);readme.append(['Purpose','Business listings for B2B contact; honour opt-outs. Public pages only.']);readme.append(['Limitations','Coverage reflects configured public sources; unverified fields remain unverified.'])
    area=wb.create_sheet('AREA SUMMARY');area.append(['Catchment','Businesses','With phone']);style_header(area)
    for a in config.areas:area.append([a, f'=COUNTIF(\'MASTER DATABASE\'!D2:D{last},A{area.max_row+1})',f'=COUNTIFS(\'MASTER DATABASE\'!D2:D{last},A{area.max_row+1},\'MASTER DATABASE\'!P2:P{last},"Yes")'])
    matrix=wb.create_sheet('GROUP x AREA');matrix.append(['Category Group',*config.areas,'Total']);style_header(matrix)
    groups=list(TREE)+['Other Local Services']
    for g in groups:
        rr=matrix.max_row+1;matrix.append([g]+[f'=COUNTIFS(\'MASTER DATABASE\'!$F$2:$F${last},$A{rr},\'MASTER DATABASE\'!$D$2:$D${last},{get_column_letter(j+2)}$1)' for j,_ in enumerate(config.areas)]+[f'=SUM(B{rr}:{get_column_letter(len(config.areas)+1)}{rr})'])
    cat=wb.create_sheet('CATEGORY TREE');cat.append(['Category Group','Subcategory','Found']);style_header(cat)
    for g,d in TREE.items():
        for sub in d['subs']:
            rr=cat.max_row+1;cat.append([g,sub,f'=COUNTIFS(\'MASTER DATABASE\'!$F$2:$F${last},A{rr},\'MASTER DATABASE\'!$G$2:$G${last},B{rr})'])
    research=wb.create_sheet('RESEARCH MATRIX');research.append(['Area','Group','Found','Target','Gap']);style_header(research)
    for a in config.areas:
        for g in groups:
            rr=research.max_row+1; research.append([a,g,f'=COUNTIFS(\'MASTER DATABASE\'!$D$2:$D${last},A{rr},\'MASTER DATABASE\'!$F$2:$F${last},B{rr})',30 if a in config.core_areas else 15,f'=MAX(0,D{rr}-C{rr})'])
    slog=wb.create_sheet('SOURCE LOG');slog.append(['Source','Term','Area','Status','Raw count','Kept count','Started','Finished','Error']);style_header(slog)
    for s in searches:slog.append([s.get(k,'') for k in ('source','term','area','status','raw_count','kept_count','started_at','finished_at','error')])
    for sheet in wb:
        if sheet.max_row>1 and sheet.auto_filter.ref is None:finish(sheet)
    output.parent.mkdir(parents=True,exist_ok=True);wb.save(output);return output
