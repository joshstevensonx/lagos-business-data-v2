from openpyxl import Workbook
from .style import *
from ..score import COMPONENTS

def build_delivery(records,config,output):
    records=sorted(records,key=lambda r:(-float(r.get('score') or 0),r.get('area',''),r.get('name','').lower()))
    wb=Workbook();ws=wb.active;ws.title='TARGET LIST'
    cols=['Rank','Business Name','Area','Category','Google Category','Delivery Score','Band','Score basis','Phone','Website','Email','Instagram','Facebook','WhatsApp','Address','Latitude','Longitude','Google Maps','Data Source','Notes']+COMPONENTS
    ws.append(cols)
    for i,r in enumerate(records,1):ws.append([i,r.get('name',''),r.get('area',''),r.get('delivery_category',''),r.get('label',''),r.get('score',''),r.get('band',''),r.get('score_basis','Floor - not enriched'),r.get('phone',''),r.get('website',''),r.get('email',''),r.get('instagram',''),r.get('facebook',''),r.get('whatsapp',''),r.get('addr',''),r.get('lat',''),r.get('lng',''),r.get('maps',''),r.get('source',''),r.get('notes','')]+[r.get('score_components',{}).get(c,0) for c in COMPONENTS])
    style_header(ws);finish(ws,freeze='C2')
    last=len(records)+1;dash=wb.create_sheet('DASHBOARD',0);dash.append(['LAGOS DELIVERY PROSPECTS']);dash.append(['KPI','Value']);style_header(dash,2)
    for band_name in ('Hot - approach first','Warm - worth a call','Cool - lower priority','Cold - little evidence of delivery'):
        dash.append([band_name,f'=COUNTIF(\'TARGET LIST\'!G2:G{last},A{dash.max_row+1})'])
    for title,formula in [('Total prospects',f'=COUNTA(\'TARGET LIST\'!B2:B{last})'),('With phone',f'=COUNTIF(\'TARGET LIST\'!I2:I{last},"<>")'),('With website',f'=COUNTIF(\'TARGET LIST\'!J2:J{last},"<>")'),('With email',f'=COUNTIF(\'TARGET LIST\'!K2:K{last},"<>")')]:dash.append([title,formula])
    readme=wb.create_sheet('README');readme.append(['ABOUT THIS LIST']);readme.append(['Score caveat','Scores are ranking aids. Floor - not enriched means delivery-page signals were not checked; absence is not evidence the business does not deliver.']);readme.append(['OSM attribution','© OpenStreetMap contributors']);readme.append(['B2B use','Contact only using published business details and honour opt-outs.'])
    method=wb.create_sheet('SCORING METHOD');method.append(['Component','Maximum points']);style_header(method)
    caps={'Google Delivery flag':25,'Catering offered':15,'Order-volume proxy (reviews)':25,'Chain / multi-branch':10,'Category weight':10,'Has website':8,'Online ordering / menu':7,'No-contact delivery':5,'Takeout':5,'Delivery hours published':5,'Rating 4.0+':5}
    for key in COMPONENTS:method.append([key,caps.get(key,0)])
    output.parent.mkdir(parents=True,exist_ok=True);wb.save(output);return output
