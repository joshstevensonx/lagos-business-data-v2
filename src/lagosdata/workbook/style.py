from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
NAVY='1F3864'; BLUE='2F5597'; LIGHT='D9E2F3'; GREY='F2F2F2'
def style_header(ws,row=1):
    for cell in ws[row]:
        cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=10); cell.fill=PatternFill('solid',fgColor=NAVY)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        cell.border=Border(*( [Side(style='thin',color='BFBFBF')]*4 ))
def finish(ws,widths=None,freeze='A2'):
    ws.freeze_panes=freeze; ws.auto_filter.ref=ws.dimensions
    if widths:
        for i,w in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=w
