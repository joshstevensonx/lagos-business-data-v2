import json, re, sys, math, collections
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter as L
sys.path.insert(0,'/home/user/lagos-business-data-v2/src'); sys.path.insert(0,'/home/user/lagos-business-data-v2/out/expand')
from lagosdata.score import score_business, band, build_name_counts
from classify_local import classify
ROOT=Path('/home/user/lagos-business-data-v2'); X=ROOT/'out/expand'
SRC=ROOT/'data/Lagos_Business_Data_MASTER_DATABASE_v4.xlsx'; OUT=sys.argv[1]
TODAY='2026-10-06'
def jl(p): return [json.loads(l) for l in open(p) if l.strip()] if Path(p).exists() else []
cands=json.load(open(X/'candidates.json')); basic=json.load(open(X/'basic_rows.json'))
pages={d['key']:d for d in jl(X/'pages.jsonl')}; crawl={d['key']:d for d in jl(X/'crawl.jsonl')}
def phone_norm(s):
    d=re.sub(r'\D','',s or '')
    if d.startswith('234') and len(d)>=13: return '+'+d[:13]
    if d.startswith('0') and len(d)==11: return '+234'+d[1:]
    return ('+'+d) if len(d)>=10 and s.strip().startswith('+') else (s or '').strip()
def info(key):
    p=pages.get(key); 
    if not p: return None
    it=p.get('items',{}); o={}
    o['address']=re.sub(r'^Address:\s*','',it.get('address',{}).get('t','')).strip()
    o['website']=(it.get('authority',{}).get('h') or '').split('?utm')[0].split('&utm')[0]
    ph=[re.sub(r'^Phone:\s*','',v['t']).strip() for k,v in it.items() if k.startswith('phone:')]
    o['phone']=phone_norm(ph[0]) if ph else ''
    od=[v for k,v in it.items() if k.startswith('action:') and re.search(r'order|deliver|menu',v['t'],re.I)]
    o['order']=re.sub(r'^.*\n','',od[0]['t']).strip() if od else ''
    o['cat']=p.get('cat','').strip(); o['h1']=p.get('h1','')
    return o
def contact(key,website):
    c=crawl.get(key,{}); s=c.get('socials',{}) or {}
    return dict(email='; '.join(c.get('emails',[])[:3]),instagram=s.get('instagram',''),facebook=s.get('facebook',''),linkedin=s.get('linkedin',''),
                wa=c.get('whatsapp',''),dsig=bool(c.get('delivery_signals')),online=bool(c.get('online_order')),phones='; '.join(c.get('phones',[])[:3]))
wb=openpyxl.load_workbook(SRC); ws=wb['MASTER DATABASE']
HDR=[c.value for c in ws[1]]; COL={h:i+1 for i,h in enumerate(HDR)}
extra=['Ordering Link (Google)','Website Delivery Wording']
for j,h in enumerate(extra):
    c=ws.cell(1,42+j,h); c._style=copy(ws['AO1']._style); ws.column_dimensions[L(42+j)].width=24
COL.update({h:42+j for j,h in enumerate(extra)})
OLD=ws.max_row
while ws.cell(OLD,2).value is None: OLD-=1
rows_by_id={ws.cell(r,1).value:r for r in range(2,OLD+1)}
BASIS='Partial - ordering link & website checked (Google delivery flags not read)'
# --- existing delivery names for chain detection
names=[{'title':ws.cell(r,2).value} for r in range(2,OLD+1) if ws.cell(r,27).value!='Community Media']
def mk(name,cat_text,reviews,rating,website,order):
    return {'name':name,'label':cat_text,'reviews':reviews or 0,'rating':rating or 0,'website':website,'order_online':bool(order)}
# --- build new records
new=[]; both=[]; skipped=collections.Counter()
for c in cands:
    if c['in_master'] and c['master'][1]!='Community Media': continue   # existing delivery row (basic ones handled below)
    inf=info(c['key'])
    if not inf: skipped['no page']+=1; continue
    cl=classify(inf['cat'],c['name'],c['terms'])
    if not cl: skipped['category: '+(inf['cat'] or 'none')]+=1; continue
    c['inf']=inf; c['cl']=cl; c['ct']=contact(c['key'],inf['website']); (both if c['in_master'] else new).append(c)
allnames=names+[{'title':c['name']} for c in new]
nc=build_name_counts(allnames)
def do_score(name,cl,reviews,rating,website,order,ct):
    rec=mk(name,cl[3],reviews,rating,website,order or ct['online'])
    rec['categoryName']=cl[3] if cl[3]!='other' else 'liquor'
    s,comp,_=score_business(rec,nc); return s,comp
def put(r,vals,style_row):
    for c in range(1,43): ws.cell(r,c)._style=copy(ws.cell(style_row,c)._style)
    for k,v in vals.items():
        if v not in (None,''): ws.cell(r,COL.get(k) or k,v)
style_row=OLD  # last delivery row
nid=max(rows_by_id)
for c in new:
    inf,cl,ct=c['inf'],c['cl'],c['ct']; nid+=1
    website=inf['website']; phone=inf['phone'] or phone_norm(c['phone']); s,comp=do_score(c['name'],cl,c['reviews'],c['rating'],website,inf['order'],ct)
    area=c['area']; zone='Core catchment' if area=='Maryland / Mende' else 'Delivery area'
    chans=sum(1 for v in (phone,website,ct['email'],ct['instagram'],ct['facebook'],ct['linkedin']) if v)
    vals={'ID':nid,'Business Name':c['name'],'Zone':zone,'Catchment':area,'Category Group':cl[1],'Subcategory':cl[2],'Source Category':inf['cat'],
          'Full Address':inf['address'],'Phone':phone,'Website':website,'Google Maps':f"https://www.google.com/maps/place/?q=place_id:{c['place_id']}" if c['place_id'].startswith('ChI') else c['maps'],
          'Google Rating':c['rating'] or None,'Google Reviews':c['reviews'] or None,'Has Phone':'Yes' if phone else 'No','Latitude':c['lat'],'Longitude':c['lng'],
          'Data Source':'Google Maps (browser sweep, Oct 2026) - place page + website check','Date Added':TODAY,'Last Checked':TODAY,'Verification':'Google-verified','Sales Status':'Not Contacted',
          'Dataset':'Delivery Prospects','Delivery Category':cl[0],'Delivery Score':s,'Delivery Band':band(s),'Delivery flag (Google)':'Not checked','Catering':'Not checked',
          'Chain':'Yes' if comp['Chain / multi-branch'] else 'No','Score basis':BASIS,'All Phones':ct['phones'] or phone,'Email':ct['email'],'Instagram':ct['instagram'],'Facebook':ct['facebook'],'LinkedIn':ct['linkedin'],'Contact Channels':chans,
          extra[0]:inf['order'],extra[1]:'Yes' if ct['dsig'] else ''}
    r=ws.max_row+1; put(r,vals,style_row)
# --- 'Both' rows: media rows that also match delivery search
for c in both:
    r=rows_by_id[c['master'][0]]; inf,cl,ct=c['inf'],c['cl'],c['ct']
    website=inf['website'] or ws.cell(r,COL['Website']).value; phone=ws.cell(r,COL['Phone']).value or inf['phone']
    s,comp=do_score(c['name'],cl,c['reviews'],c['rating'],website,inf['order'],ct)
    chans=sum(1 for v in (phone,website,ct['email'],ct['instagram'],ct['facebook'],ct['linkedin']) if v)
    for k,v in {'Dataset':'Both','Delivery Category':cl[0],'Delivery Score':s,'Delivery Band':band(s),'Delivery flag (Google)':'Not checked','Catering':'Not checked',
                'Chain':'Yes' if comp['Chain / multi-branch'] else 'No','Score basis':BASIS,'Email':ct['email'],'Instagram':ct['instagram'],'Facebook':ct['facebook'],'LinkedIn':ct['linkedin'],'Contact Channels':chans,extra[0]:inf['order'],extra[1]:'Yes' if ct['dsig'] else ''}.items():
        if v not in (None,''): ws.cell(r,COL.get(k) or k,v)
    if not ws.cell(r,COL['Website']).value and website: ws.cell(r,COL['Website'],website)
# --- re-enrich the 232 Basic rows
upd=0
for b in basic:
    r=rows_by_id[b['master_id']]; inf=info(b['key'])
    if not inf: continue
    ct=contact(b['key'],inf['website']); name=ws.cell(r,2).value
    website=ws.cell(r,COL['Website']).value or inf['website']; phone=ws.cell(r,COL['Phone']).value or inf['phone']
    reviews=ws.cell(r,COL['Google Reviews']).value; rating=ws.cell(r,COL['Google Rating']).value
    dc=ws.cell(r,COL['Delivery Category']).value
    cl=('', '', '', {'Pharmacy':'pharmacy','Supermarket / Grocery':'supermarket'}.get(dc,'restaurant'))
    s,comp=do_score(name,cl,reviews,rating,website,inf['order'],ct)
    chans=sum(1 for v in (phone,website,ct['email'],ct['instagram'],ct['facebook'],ct['linkedin']) if v)
    def setv(k,v): ws.cell(r,COL.get(k) or k,v)
    if website and not ws.cell(r,COL['Website']).value: setv('Website',website)
    if phone and not ws.cell(r,COL['Phone']).value: setv('Phone',phone); setv('Has Phone','Yes')
    if inf['address'] and not ws.cell(r,COL['Full Address']).value: setv('Full Address',inf['address'])
    for k,v in {'Delivery Score':s,'Delivery Band':band(s),'Delivery flag (Google)':'Not checked','Catering':'Not checked','Chain':'Yes' if comp['Chain / multi-branch'] else 'No','Score basis':BASIS,
                'Email':ct['email'],'Instagram':ct['instagram'],'Facebook':ct['facebook'],'LinkedIn':ct['linkedin'],'Contact Channels':chans,extra[0]:inf['order'],extra[1]:'Yes' if ct['dsig'] else ''}.items():
        if v not in (None,''): setv(k,v)
    setv('Last Checked',TODAY); upd+=1
NEW=ws.max_row
# --- re-rank all delivery rows
drows=[r for r in range(2,NEW+1) if ws.cell(r,27).value!='Community Media']
drows.sort(key=lambda r:(-(ws.cell(r,30).value or 0),-(ws.cell(r,15).value or 0)))
for k,r in enumerate(drows,1): ws.cell(r,29,k)
print('new rows',len(new),'both',len(both),'basic updated',upd,'skipped',dict(skipped),'last row',NEW)
json.dump({'NEW':NEW,'OLD':OLD,'new':len(new),'both':len(both),'upd':upd,'skipped':dict(skipped)},open(X/'stats.json','w'))
wb.save(OUT)
