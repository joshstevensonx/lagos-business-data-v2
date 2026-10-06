import sys, re, json
from copy import copy
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.cell_range import MultiCellRange
IN,OUT=sys.argv[1:3]
st=json.load(open('/home/user/lagos-business-data-v2/out/expand/stats.json')); NEW=st['NEW']; OLD=5721
wb=openpyxl.load_workbook(IN); ws=wb['MASTER DATABASE']
M="'MASTER DATABASE'!"
def R(col): return f"{M}${col}$2:${col}${NEW}"
ws.auto_filter.ref=f'A1:AQ{NEW}'
for dv in ws.data_validations.dataValidation:
    dv.sqref=MultiCellRange(re.sub(str(OLD),str(NEW),str(dv.sqref)))
    if dv.formula1 and 'Anthony / Anthony Village' in dv.formula1 and 'Ikeja' not in dv.formula1:
        dv.formula1=dv.formula1[:-1]+',Ikeja,Ajah,Apapa,Lekki Phase 2"'
for s in wb.sheetnames:
    if s in ('MASTER DATABASE',): continue
    for row in wb[s].iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('=') and str(OLD) in c.value: c.value=c.value.replace(str(OLD),str(NEW))
NEWAREAS=['Ikeja','Ajah','Apapa','Lekki Phase 2']
# ---- AREA SUMMARY: insert 4 rows before total
a=wb['AREA SUMMARY']; TOT0=17
body={c:copy(a.cell(16,c)._style) for c in range(1,12)}; tot={c:copy(a.cell(TOT0,c)._style) for c in range(1,12)}
for c in range(1,12): a.cell(TOT0,c).value=None
allareas=[a.cell(r,1).value for r in range(4,TOT0)]+NEWAREAS
TOT=4+len(allareas)
for k,n in enumerate(allareas):
    r=4+k
    if n in NEWAREAS:
        a.cell(r,1,n); a.cell(r,2,'Delivery area')
        f={3:f'=COUNTIF({R("D")},$A{r})',4:f'=IFERROR(C{r}/$C${TOT},0)',5:f'=COUNTIFS({R("D")},$A{r},{R("P")},"Yes")',6:f'=COUNTIFS({R("D")},$A{r},{R("V")},"Google-verified")',
           7:f'=COUNTIFS({R("D")},$A{r},{R("W")},"A - High commercial relevance")',8:f'=COUNTIFS({R("D")},$A{r},{R("W")},"B - Potential advertiser")',
           9:f'=COUNTIFS({R("D")},$A{r},{R("W")},"C - Ring area, contactable")+COUNTIFS({R("D")},$A{r},{R("W")},"D - Listing only (no contact yet)")',
           10:f'=COUNTIFS({R("D")},$A{r},{R("AE")},"Hot - approach first")',11:f'=COUNTIFS({R("D")},$A{r},{R("AE")},"Warm - worth a call")'}
        for c in range(1,12):
            a.cell(r,c)._style=copy(body[c])
            if c in f: a.cell(r,c,f[c])
    else:
        for c in (4,): a.cell(r,c,f'=IFERROR(C{r}/$C${TOT},0)')
a.cell(TOT,1,'TOTAL')
for c in range(1,12):
    a.cell(TOT,c)._style=copy(tot[c])
    if c>=3: a.cell(TOT,c,f'=SUM({L(c)}4:{L(c)}{TOT-1})')
# existing rows: refresh Delivery Hot/Warm for all (Maryland etc.) already formulas; ok
# Priority C/D for Maryland etc unchanged
# ---- GROUP x AREA: add 4 cols before Total
g=wb['GROUP x AREA']; mr=g.max_row; TC=15
for r in range(3,mr+1):
    tv=g.cell(r,TC).value; ts=copy(g.cell(r,TC)._style); bs=copy(g.cell(r,TC-1)._style)
    for k,n in enumerate(NEWAREAS):
        c=g.cell(r,TC+k); c._style=copy(bs)
        if r==3: c.value=n
        else: c.value=Translator(g.cell(r,TC-1).value,origin=f'{L(TC-1)}{r}').translate_formula(f'{L(TC+k)}{r}')
    o=g.cell(r,TC+4); o._style=ts
    if r==3: o.value=tv
    else: o.value=re.sub(r'B(\d+):N(\d+)',r'B\1:R\2',tv) if 'SUM(B' in tv else tv.replace('O','S')
for k in range(5): g.column_dimensions[L(TC+k)].width=g.column_dimensions['N'].width or 14
g.merged_cells.ranges=set(); g.merge_cells('A1:S1')
g['A2']='Anthony Community Media catchments (B-H) and delivery-prospect areas (I-R).'
# ---- DELIVERY SUMMARY
d=wb['DELIVERY SUMMARY']
DL=f'{R("AA")},"<>Community Media"'
d['A2']='Restaurants, supermarkets, pharmacies, bakeries, cafes, caterers and drinks shops across 11 Lagos areas, ranked by how likely they are to need a delivery partner.'
hl={5:f'=COUNTIFS({DL})',6:f'=COUNTIFS({DL},{R("AE")},"Hot - approach first")',7:f'=COUNTIFS({DL},{R("AE")},"Warm - worth a call")',8:f'=COUNTIFS({DL},{R("AF")},"Yes")',9:f'=COUNTIFS({DL},{R("AG")},"Yes")',
    10:f'=COUNTIFS({DL},{R("AH")},"Yes")',11:f'=COUNTIFS({DL},{R("P")},"Yes")',12:f'=COUNTIFS({DL},{R("L")},"<>")',13:f'=COUNTIFS({DL},{R("AK")},"<>")',14:f'=COUNTIFS({DL},{R("AL")},"<>")',
    15:f'=COUNTIFS({DL},{R("AI")},"Full - delivery data checked")'}
for r,f in hl.items(): d.cell(r,3,f)
for col in 'ABCDEF': d[f'{col}16']._style=copy(d[f'{col}15']._style)
d['A16']='Partly checked (ordering link + website, no Google flags)'; d['C16']=f'=COUNTIFS({DL},{R("AI")},"Partial*")'
hdr=copy(d['B18']._style); body=copy(d['B19']._style); bodyA=copy(d['A19']._style); totA=copy(d['A25']._style); totB=copy(d['B25']._style)
cats=['Restaurant / Food','Supermarket / Grocery','Pharmacy','Bakery / Desserts','Cafe / Coffee','Catering','Drinks / Liquor']
areas=['Lekki Phase 1','Lekki / Ajah corridor','Victoria Island','Ikoyi','Yaba','Surulere','Ikeja','Ajah','Apapa','Lekki Phase 2','Maryland / Mende']
for r in range(18,40):
    for c in range(1,13): d.cell(r,c).value=None
d.cell(18,1,'Area')._style=copy(d['A18']._style)
heads=cats+['Total','Hot','Warm']
for j,h in enumerate(heads): d.cell(18,2+j,h)._style=copy(hdr)
for k,n in enumerate(areas):
    r=19+k; d.cell(r,1,n)._style=copy(bodyA)
    for j,cn in enumerate(cats): d.cell(r,2+j,f'=COUNTIFS({R("D")},$A{r},{R("AB")},"{cn}")')._style=copy(body)
    d.cell(r,9,f'=SUM(B{r}:H{r})')._style=copy(body)
    d.cell(r,10,f'=COUNTIFS({R("D")},$A{r},{R("AE")},"Hot - approach first")')._style=copy(body)
    d.cell(r,11,f'=COUNTIFS({R("D")},$A{r},{R("AE")},"Warm - worth a call")')._style=copy(body)
RR=19+len(areas); d.cell(RR,1,'Other Community Media catchments (tagged Both)')._style=copy(bodyA)
for j,cn in enumerate(cats): d.cell(RR,2+j,f'=COUNTIFS({DL},{R("AB")},"{cn}")-SUM({L(2+j)}19:{L(2+j)}{RR-1})')._style=copy(body)
d.cell(RR,9,f'=SUM(B{RR}:H{RR})')._style=copy(body)
d.cell(RR,10,f'=COUNTIFS({DL},{R("AE")},"Hot - approach first")-SUM(J19:J{RR-1})')._style=copy(body)
d.cell(RR,11,f'=COUNTIFS({DL},{R("AE")},"Warm - worth a call")-SUM(K19:K{RR-1})')._style=copy(body)
TR=RR+1; d.cell(TR,1,'TOTAL')._style=copy(totA)
for c in range(2,12): d.cell(TR,c,f'=SUM({L(c)}19:{L(c)}{TR-1})')._style=copy(totB)
for c in range(2,12): d.column_dimensions[L(c)].width=max(d.column_dimensions[L(c)].width or 12,14)
d.merged_cells.ranges=set(); d.merge_cells('A1:K1'); d.merge_cells('A2:K2')
# ---- DASHBOARD
db=wb['DASHBOARD']
db['A2']='Anthony Community Media catchments (7)  |  Delivery-prospect areas (11)  |  29 category groups  |  Last refreshed 2026-10-06'
db['A13']='Delivery-prospect areas (Lekki, VI, Ikoyi, Yaba, Surulere, Ikeja, Ajah, Apapa)'
# ---- SOURCE LOG
sl=wb['SOURCE LOG']; r0=sl.max_row
add=[('Google Maps (browser sweep, Oct 2026)',f"589 free browser searches (19 terms x 11 areas x 2-3 viewports) using the repo's own collector. Added {st['new']:,} new businesses and delivery fields for {st['both']} existing Community Media rows.",st['new']+st['both'],'Delivery Prospects expansion: discovery, then place-page check for address, website, phone and ordering link','https://www.google.com/maps'),
     ('Google Maps place pages + business websites (free)',f"Re-checked {st['upd']} previously 'Basic' prospects for website, phone, ordering link, emails and social handles (robots-aware crawl). Google's delivery/catering flags are NOT visible in the free signed-out view.",st['upd'],'Re-scored with the same weights; Score basis shows Partial','https://www.google.com/maps')]
for k,vals in enumerate(add):
    for c,v in enumerate(vals,1):
        cell=sl.cell(r0+1+k,c,v); cell._style=copy(sl.cell(r0,c)._style)
# ---- DELIVERY SCORING note
ds=wb['DELIVERY SCORING']; n=ds.max_row+2
ds.cell(n,1,'NOTE ON THE COUNTS COLUMN').font=copy(ds['A17'].font)
ds.cell(n+1,1,f"Counts in column D describe the original 774-prospect list. Rows added on 2026-10-06 ({st['new']:,} new + {st['upd']} re-checked) were scored with the same weights but without the Google delivery, no-contact delivery, takeout, delivery-hours and catering signals (not visible in the free browser view). Their Score basis reads 'Partial'; treat their score as a floor.")
ds.cell(n+1,1).alignment=openpyxl.styles.Alignment(wrap_text=True,vertical='top'); ds.merge_cells(start_row=n+1,start_column=1,end_row=n+1,end_column=4); ds.row_dimensions[n+1].height=75
# ---- README
rd=wb['README']
rd['B3']='One master list combining the Anthony Community Media business directory (7 catchments) and the Lagos Delivery Prospects list (11 areas: restaurants, supermarkets, pharmacies, bakeries, cafes, caterers and drinks shops). Each record keeps its original fields; the Dataset column (AA) says which list it came from (Community Media / Delivery Prospects / Both).'
for r in range(1,rd.max_row+1):
    if rd.cell(r,1).value=='Delivery areas':
        rd.cell(r,2,'Delivery rows have Zone = "Delivery area" and Catchment = the area (Lekki Phase 1, Lekki / Ajah corridor, Victoria Island, Ikoyi, Yaba, Surulere, Ikeja, Ajah, Apapa, Lekki Phase 2). A few Maryland / Mende delivery finds sit in Core catchment. Delivery rows are classified into the 29-group tree; the original delivery category (Restaurant / Food, Supermarket / Grocery, Pharmacy, Bakery / Desserts, Cafe / Coffee, Catering, Drinks / Liquor) is kept in Delivery Category (AB). Prospect Priority and Suggested Package are blank for them - use Delivery Score / Band instead.')
    if rd.cell(r,1).value=='Merge rules':
        rd.cell(r,2,f"No Community Media records were dropped. {st['both']} Community Media businesses that also appeared in the delivery sweep were tagged Dataset = Both and given delivery fields. Duplicates were matched on Google place ID, then on name plus location within about 60 m. Sales Status \"Not contacted\"/\"Not approached\" was standardised to \"Not Contacted\". IDs 5721 onward are the 2026-10-06 expansion.")
    if rd.cell(r,1).value=='Delivery caveat':
        rd.cell(r,2,"Only the first 542 delivery prospects were fully checked for Google's delivery/takeout/catering flags (Score basis = Full). Every other delivery row, including all rows added on 2026-10-06, was scored from the ordering link, website, reviews, chain status, category and rating, so its score is a floor. 'Not checked' in the Delivery flag and Catering columns means unknown, not No. See DELIVERY SCORING.")
rn=rd.max_row+1
for c,v in ((1,'Expansion method'),(2,"Free tools only: the repo's own Google Maps browser collector (no paid APIs, no Apify, no paid scraping), plus a robots-aware crawl of public business websites. Runs were rate-limited and would have halted on any CAPTCHA/consent wall (none occurred).")):
    cell=rd.cell(rn,c,v); cell._style=copy(rd.cell(rn-1,c)._style)
rd.row_dimensions[rn].height=rd.row_dimensions[rn-1].height
wb.save(OUT)
print('saved',NEW)
